"""Read the supplied schema; create QMD review tables. Never writes the XLSX.

Run only before manual review edits: archived inventory pages are replaced; published appendices are not modified.
Colleague decisions live separately in Redaktion/Archiv/preset-vollinventar/preset-review.md.
"""
from pathlib import Path
import hashlib
import json
from collections import defaultdict
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'Redaktion/Primärquellen/schema_1.13.1_dev.xlsx'
wb = openpyxl.load_workbook(SOURCE, data_only=False)
inp = defaultdict(list)
for row in wb['IN'].iter_rows(min_row=2):
    if row[7].value:
        inp[str(row[7].value)].append((row[0].row, row[9].value, row[10].value, row[12].value))

# Editorial ordering, independent of the export's physical sheet order.
GROUPS = [
 ('quartier', 'Quartier, Geometrie und Bewertungsumfang', [
  ('renovation','Sanierungsanteile','Sanierungsflächen und Neubauanteile je Nutzung abgrenzen.'),
  ('geometry','Gebäudetypen und Bauteilmengen','Typologie, Bauperiode, Flächenbezug und Mengengerüst gemeinsam prüfen. „Aktuelle Hülle“ ist ein Projektbezug, kein allgemeiner Default.'),
  ('declaration_type','Bilanzumfang der Qualitätsstufen','Der enthaltene Mobilitätsschalter dokumentiert nicht allein alle Anforderungen der drei Qualitätsstufen.')]),
 ('nutzung','Nutzung und Energiebedarf',[
  ('plugloads','Nutzerstrom','Jahresbezug, NGF/BGF, innere Gewinne und Abgrenzung zu Allgemeinstrom/Beleuchtung prüfen. „Kein Nutzerstrom“ ist im Schema für Studienzwecke bezeichnet.'),
  ('dhw_demand','Warmwasserwärmebedarf','Bedarf je Nutzung und Flächenbezug begründen. Die beiden OIB-Presetspalten sind im Export noch leer; ihr Name ist kein Normbeleg.')]),
 ('huelle','Gebäudehülle und sommerlicher Wärmeschutz',[
  ('hulls','Wärmeschutz der Gebäudehülle','U-Werte, g-Werte und Wärmebrückenzuschlag gemeinsam nach Bauqualität und Gültigkeit prüfen.'),
  ('air_tightness','Luftdichtheit','n50, Bauzustand und Nachweisannahme prüfen; ähnlich benannte Presets auf Redundanz prüfen.'),
  ('building_mass','Wirksame Speichermasse','Wirksame Wärmekapazität und Flächenbezug prüfen; keine Gleichsetzung mit der gesamten Bauteilmasse.'),
  ('shading_geometry','Geometrische Einstrahlungsfaktoren','Orientierung, Sommer/Winter, Rahmenanteil und Bezugsnorm klären.'),
  ('mobile_shading','Beweglicher Sonnenschutz','Fc-Werte, Verfügbarkeit und Bedienannahme je Nutzung prüfen.')]),
 ('haustechnik','Heizung, Kühlung, Warmwasser und Lüftung',[
  ('heating_system','Wärmeerzeugung','Leistungsbezug, Wirkungsgrad/COP, Hilfsenergie, Verteilung und Systemtemperaturen prüfen.'),
  ('emission_system','Wärme- und Kälteabgabe','Heiz- und Kühlleistung der Abgabesysteme einschließlich Komfort- und Temperaturgrenzen prüfen.'),
  ('heating_period','Heizperioden','Monatliche Freigabe und Ganzjahresbetrieb von tatsächlich auftretendem Heizbedarf unterscheiden.'),
  ('cooling_system','Kälteerzeugung','Freecooling, aktive Kälte, Leistungsbezüge und Hilfsenergie getrennt prüfen.'),
  ('cooling_modes','Anteile aktiv gekühlter Nutzungen','Gekühlte Flächenanteile mit dem Nutzungskonzept und Zwei-Zonen-Modell abgleichen.'),
  ('cooling_period','Kühlperioden','Monatliche Freigabe, Komfortanforderung und tatsächlichen Bedarf unterscheiden.'),
  ('dhw_system','Warmwasserbereitung und Verteilung','Systemanteile, Speicher- und Verteilverluste sowie Temperaturniveau gemeinsam prüfen.'),
  ('ventilation_system','Lüftungstechnik und Wärmerückgewinnung','Lüfterstrom-Skalierung, Wirkungsgrade und Sommerbetrieb prüfen; Skalierungsfaktor ist kein absoluter Strombedarf.'),
  ('vent_modes','Anteile mechanisch belüfteter Nutzungen','Anteile je Nutzung und hygienisch erforderlichen Luftwechsel zusammen prüfen.'),
  ('vent_night_modes','Nachtauskühlung','Fensterstellung und Verfügbarkeit prüfen; Lärm, Sicherheit und Regelung als Anwendungsvoraussetzungen benennen.')]),
 ('energie','Photovoltaik, Speicher und Flexibilität',[
  ('PV_system','Photovoltaikbelegung','Bezugsfläche und Bedeutung von Belegungswerten über 100 % klären; Flächenpreset ersetzt kein Ertragsprofil.'),
  ('battery_system','Batteriespeicher','Kapazitätsableitung aus PV-Leistung, C-Rate und Lade-/Entladewirkungsgrad prüfen. Presetbezeichnung und exportierter Kapazitätswert können unterschiedliche Ebenen beschreiben.'),
  ('flexibility_system','Nutzung erneuerbarer Überschüsse','Freigaben für Nutzerstrom, Mindestversorgung und Fahrzeuge mit Signaldefinition und Bewertung abgleichen.'),
  ('flexibility_thermal','Thermische Flexibilität','Temperaturhub, gewähltes System und Komfortgrenzen prüfen.')]),
 ('mobilitaet','Mobilität und Standort',[
  ('mobility','Mobilitätsregionen','Im Preset wird eine Gemeinde ausgewählt. Fahrleistung, Modal Split und Flotte sind dadurch noch nicht als vollständiger Datensatz dokumentiert.')]),
 ('bauweise','Bauweisen und Lebenszyklus',[
  ('construction_type','Bauweisen und Bauteilzuordnungen','Bauteilnamen sind Verweise, keine THG-Kennwerte. Datensatz, funktionelle Einheit, Module, Nutzungsdauer und biogene Bilanzierung ergänzen. Live-/Testspalten vor Freigabe gesondert entscheiden.')])]

def val(v):
    if v is None: return '—'
    if isinstance(v,bool): return 'ja' if v else 'nein'
    if isinstance(v,float): return format(v,'.8g').replace('.',',')
    return str(v)

def esc(v):
    return val(v).replace('|',r'\|').replace('\r',' ').replace('\n',' / ').replace('<','&lt;').replace('>','&gt;')

def code(v): return '`'+str(v).replace('`','')+'`' if v is not None else '—'
def table(headers, rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(str(x) for x in r)+' |' for r in rows])+'\n\n'

inventory=[]
coverage=[]
issues=[]
pages=[]
tid=0
for slug,title,items in GROUPS:
    filename=f'anhang-presets-{slug}.qmd'
    pages.append(filename)
    out=[f'# {title} {{#sec-presets-{slug}}}\n', '::: {.preset-anhang}\n',
         'Diese Tabellen sind eine **Prüffassung aus dem Entwicklungsschema 1.13.1_dev**, keine freigegebenen Standardwerte. Leseschlüssel und Rückmeldeverfahren stehen in @sec-preset-leseschluessel.\n']
    for sheet,title,note in items:
        tid+=1; key=f'T{tid:02d}'; s=wb[sheet]
        special=sheet in ('geometry','construction_type')
        if sheet=='geometry':
            ridx=list(range(3,s.max_row+1)); cols=list(range(5,s.max_column+1)); varcol=3; labelcol=1; unitcol=2
        elif sheet=='construction_type':
            ridx=list(range(3,s.max_row+1));cols=[4]+list(range(6,s.max_column+1));varcol=3;labelcol=1;unitcol=None
        else:
            ridx=[r for r in range(4,s.max_row+1) if s.cell(r,1).value not in (None,'unit')]
            cols=list(range(4,s.max_column+1));varcol=1;labelcol=2;unitcol=3
        # Empty trailing formatted columns are not new presets.
        cols=[c for c in cols if any(s.cell(r,c).value is not None for r in range(1,s.max_row+1))]
        selector='preset_recorded_'+('heating_emission_system' if sheet=='emission_system' else sheet)
        selector_text=code(selector) if selector in inp else 'Kein entsprechender Auswahlparameter in IN eindeutig gefunden.'
        out += [f'## {key}: {title} {{#sec-preset-{key.lower()}}}\n',note+'\n',
                f'**PEExcel-Bezug:** Schemablatt {code(sheet)}, Bereich {code(s.calculate_dimension())}. Auswahlparameter: {selector_text}\n',
                '### Parameter und fachliche Quellen\n']
        prows=[]
        labels={}
        for r in ridx:
            var=s.cell(r,varcol).value; label=s.cell(r,labelcol).value
            found=inp.get(str(var),[]) if var is not None else []
            if not label or str(label).startswith('#'):
                if len(found)==1 and found[0][1]:label=found[0][1]
                else:label='Bezeichnung ergänzen'
            unit=s.cell(r,unitcol).value if unitcol else None
            unit_source='Preset'
            if sheet=='plugloads':
                unit='kWh/m²NGFa';unit_source='Preset-Zeile 4'
            if not unit:
                units={str(x[2]) for x in found if x[2] is not None}
                if len(units)==1:unit=units.pop();unit_source='IN'
                elif sheet=='plugloads':unit='kWh/m²NGFa';unit_source='Preset-Zeile 4'
                else:unit='offen';unit_source=''
            loc=', '.join('IN!H'+str(x[0]) for x in found) or ('kein IN-Treffer' if var else 'Metadaten/Zuordnung offen')
            if len(found)>1:issues.append(f'{key}.R{r:02d}: mehrere IN-Treffer für {var}: {loc}')
            if str(unit)=='offen':issues.append(f'{key}.R{r:02d}: Einheit fehlt ({var or label})')
            pkey=f'{key}.R{r:02d}'
            labels[r]=label
            prows.append([pkey,esc(label)+' / '+code(var),esc(unit)+(f' ({unit_source})' if unit_source else ''),esc(loc),'zu ergänzen'])
        out.append(table(['Zeile','Parameter / PEExcel-Kennung','Einheit laut Schema','Fundstelle','Fachquelle / Gültigkeit'],prows))
        out += ['### Presets und Varianten\n']
        crows=[]
        for c in cols:
            col=get_column_letter(c);header=s.cell(1,c).value;display=s.cell(2,c).value
            role='Preset, ungeprüft'
            if str(header).lower()=='live':role='Live-Projektspalte, kein Default'
            elif str(header).lower() in ('default','auswählen') or str(display).lower()=='auswählen':role='Auswahl-/Defaultspalte, Freigabe offen'
            elif header=='Aktuelle Hülle':role='Projektabhängige Ableitung'
            elif 'test' in str(header).lower() or header=='Spalte2':role='Test-/Hilfsspalte, Verwendung klären'
            desc=s.cell(3,c).value if not special else None
            crows.append([key+'.P'+col,esc(header),esc(display),role,esc(desc)])
        out.append(table(['Preset-ID / Excel-Spalte','Spaltenname im Schema','Anzeige','Einordnung','Anwendung laut Export'],crows))
        out += ['### Werte zur fachlichen Prüfung\n',
                'Zeilen-IDs verweisen auf den Parameterkatalog oben; Preset-IDs auf die Excel-Spalten. **— bedeutet leer im Export, nicht null.** Zahlen bleiben in der ausgewiesenen Schemaeinheit; Verhältniswerte werden nicht stillschweigend zu Prozentwerten umgerechnet.\n']
        if sheet=='geometry':
            out += ['Die umfangreichen Typologiewerte sind zum gezielten Nachschlagen aufklappbar. Für die erste Prüfrunde genügen Typenabdeckung, Parameterkatalog und die für den jeweiligen Fachbereich relevanten Typologien.\n',
                    '::: {.callout-note collapse="true"}\n## Vollständige Geometriewerte\n']
        ncols=2 if sheet=='construction_type' else 4
        for chunkstart in range(0,len(cols),ncols):
            chunk=cols[chunkstart:chunkstart+ncols]
            headers=['Parameter']+[key+'.P'+get_column_letter(c)+': '+esc(s.cell(1,c).value) for c in chunk]
            vrows=[]
            for r in ridx:
                vrows.append([f'{key}.R{r:02d}: '+esc(labels[r])]+[esc(s.cell(r,c).value) for c in chunk])
                for c in chunk:
                    coverage.append((sheet,s.cell(r,c).coordinate))
            out.append(table(headers,vrows))
        if sheet=='geometry':out.append(':::\n')
        out += ['### Festlegung nach Prüfung\n',table(['Prüffrage','Rückmeldung'],[
            ['Fehlende / überflüssige Presets oder Parameter','offen'],
            ['Empfohlener Standard und zulässige Alternativen','offen'],
            ['Gültigkeit, Quellen und Anwendungsvoraussetzungen','offen'],
            ['Freigabe durch / Datum','ausständig']])]
        inventory.append({'id':key,'sheet':sheet,'title':title,'file':filename,'rows':len(ridx),'columns':len(cols)})
    out.append(':::\n')
    (ROOT/'Redaktion/Archiv/preset-vollinventar'/filename).write_text('\n'.join(out),encoding='utf-8')

nonpresets={'META','IN','OUT','CHART_metadata','CHART_types','SIM2'}
assert {v['sheet'] for v in inventory}==set(wb.sheetnames)-nonpresets
assert len(coverage)==len(set(coverage))
expected=sum(v['rows']*v['columns'] for v in inventory)
assert len(coverage)==expected
manifest={'source':str(SOURCE.relative_to(ROOT)),'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
          'schema_version':wb['META']['B4'].value,'inventory':inventory,'covered_cells':len(coverage),
          'issues':issues,'pages':pages}
(ROOT/'Redaktion/preset-inventar.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'{len(inventory)} Presetgruppen, {len(coverage)} Wertefelder, {len(pages)} Anhangsseiten; XLSX unverändert.')
