"""Technical I/O and display helpers only; scientific calculations live in the notebook."""
from pathlib import Path
import hashlib
import pandas as pd
import openpyxl

def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def read_source(path):
    """Read cached Excel values, preserving blanks and error strings."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    def frame(name, header=0, max_col=None):
        rows = list(book[name].iter_rows(values_only=True, max_col=max_col))
        return pd.DataFrame(rows[header+1:], columns=rows[header]).dropna(how='all')
    tables = {name:frame(name) for name in ['Strombed_estim','reduced','Building-Sector_Pool_total','save']}
    tables['zuordnung_gfz'] = frame('zuordnung_gfz', max_col=6)
    tables['simulation'] = frame('Tabelle2', header=3)
    meta = list(book['Tabelle2'].iter_rows(min_row=1,max_row=4,values_only=True))
    tables['dictionary'] = pd.DataFrame({'metric':meta[3], 'label':meta[0], 'unit':meta[1], 'note':meta[2]})
    book.close()
    return tables

def markdown_table(frame, caption=None, identifier=None):
    """Format a small table without an optional tabulate dependency."""
    def fmt(v):
        if pd.isna(v): return 'offen'
        if isinstance(v,float): return f'{v:,.3f}'.replace(',','X').replace('.',',').replace('X','.')
        return str(v).replace('|',' / ').replace('\n',' ')
    result = '| '+' | '.join(map(str,frame.columns))+' |\n'
    result += '| '+' | '.join(['---']*len(frame.columns))+' |\n'
    result += '\n'.join('| '+' | '.join(fmt(v) for v in row)+' |' for row in frame.itertuples(index=False,name=None))
    if caption: result += '\n\n: '+caption+(f' {{#{identifier}}}' if identifier else '')+'\n'
    return result
